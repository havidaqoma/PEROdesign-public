"""Paper A Source Data workbook: every raw call, the data behind every figure and table, every library value and assumption.
  python scripts/paperA_sourcedata.py   -> paper/paperA/build/PaperA_SourceData.xlsx
All cells are written from the stored files; the script re-reads the workbook afterwards and checks row counts and a
sample of values against the run files (fails loudly on any mismatch)."""
import json
import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from perodesign.library import load_library  # noqa: E402
from perodesign.optim import objective as O  # noqa: E402
from perodesign.optim import search as S  # noqa: E402
from perodesign.tea import load_cost_library  # noqa: E402

A = REPO / "phase0" / "paperA"
P = REPO / "paper" / "paperA"
OUT = P / "build" / "PaperA_SourceData.xlsx"
# --merged: the copy shipped with the merged manuscript (paper/paperAB), named by content
MERGED = "--merged" in sys.argv
if MERGED:
    OUT = REPO / "paper" / "paperAB" / "build" / "SourceData_thickness_only.xlsx"
ARCHS, SCENS, METHODS = ("p-i-n", "n-i-p"), ("base", "worse", "better"), ("random", "nsga2", "qnehvi")
# main-text numbering of the merged manuscript (paperAB build order: tab:space, tab:gap, tab:designs, tab:bgap,
# tab:q1, tab:speed, tab:bq1; fig:env, fig:fronts, fig:parts, fig:bfronts, fig:bzoom, fig:hv)
MERGED_TABLES = {1: 1, 2: 5, 3: 6, 4: 2, 5: 3}
MERGED_FIGS = {2: 6, 3: 2, 4: 3}


def merged_numbering(wb):
    """Renumber main-text table/figure references (sheet names and text cells) to the merged manuscript."""
    import re as _re

    def fix(text):
        # only main-text references ("main Table 2", "Main Figure 3") and sheet names ("Table2_", "Fig3_");
        # literature citations such as "Shahinuzzaman 2022 Table 2" are left untouched
        text = _re.sub(r"\b([Mm]ain Table )([1-5])\b",
                       lambda m: f"{m.group(1)}{MERGED_TABLES[int(m.group(2))]}", text)
        text = _re.sub(r"\b([Mm]ain Figure )([2-4])\b",
                       lambda m: f"{m.group(1)}{MERGED_FIGS[int(m.group(2))]}", text)
        text = _re.sub(r"^(\s*Table)([1-5])_", lambda m: f"{m.group(1)}{MERGED_TABLES[int(m.group(2))]}_", text)
        return _re.sub(r"^(\s*Fig)([2-4])_", lambda m: f"{m.group(1)}{MERGED_FIGS[int(m.group(2))]}_", text)
    for ws in wb.worksheets:
        if ws.title != "Calls_all":
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and ("Table" in c.value or "Fig" in c.value):
                        c.value = fix(c.value)
        ws.title = fix(ws.title)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="305496")
NOTE = Font(italic=True, color="555555")


def sheet(wb, name, title, cols, rows, widths=None, notes=()):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    for n in notes:
        ws.append([n])
        ws.cell(ws.max_row, 1).font = NOTE
    ws.append([])
    ws.append(cols)
    hr = ws.max_row
    for j in range(1, len(cols) + 1):
        c = ws.cell(hr, j)
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([v if not isinstance(v, (list, tuple, dict)) else json.dumps(v) for v in r])
    ws.freeze_panes = ws.cell(hr + 1, 1)
    for j, col in enumerate(cols, 1):
        w = (widths or {}).get(col) or min(60, max(10, len(str(col)) + 2))
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws, hr


def load_runs():
    runs = {}
    for p in sorted((A / "runs").glob("*.json")):
        r = json.loads(p.read_text())
        runs[(r["arch"], r["scenario"], r["method"], r["seed"])] = r
    assert len(runs) == 90
    return runs


def main():
    runs = load_runs()
    res = json.loads((A / "result.json").read_text())
    si = json.loads((A / "si" / "si_result.json").read_text())
    p2 = json.loads((REPO / "phase0" / "sensitivity" / "result.json").read_text())
    N = json.loads((P / "numbers.json").read_text(encoding="utf-8"))
    NS = json.loads((P / "numbers_si.json").read_text(encoding="utf-8")) if (P / "numbers_si.json").exists() else {}
    lib, clib = load_library(), load_cost_library()
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------------ README
    ws = wb.create_sheet("README")
    lines = [
        ("Source Data for: Cost-aware design of perovskite solar modules with drift-diffusion simulation and "
         "multi-objective optimisation. Thickness-only design: the optimiser sets layer thicknesses and stripe width, "
         "materials fixed (companion workbook: SourceData_material_choice.xlsx)") if MERGED else
        "Source Data for: A drift-diffusion to factory-cost environment for perovskite module design (Paper A)",
        "Author: Havid Aqoma. Generated by scripts/paperA_sourcedata.py from the stored run files; no value typed by hand.",
        "",
        "Status labels used in this workbook:",
        "  measured   = read from a simulator call record (phase0/paperA/runs/*.json)",
        "  computed   = derived by a script from measured values (formula or source named in the sheet)",
        "  sourced    = a published value; the cost library entry carries its verbatim quotation and source-file SHA-256",
        "  assumed    = a value chosen by the authors with no source; tagged 'assumed' everywhere it is used",
        "  post-hoc   = analysis added after the pre-registered analysis (Supplementary Notes S9, S10)",
        "All efficiencies and costs are model predictions (no device validation).",
        "",
        "Sheets:",
    ]
    sheets_doc = [
        ("Calls_all", "every simulator call of all 90 runs (21,600 rows): inputs, simulator outputs, module factors, objectives"),
        ("Runs_summary", "one row per run: final hypervolume, infeasible calls, calls to 95%, wall time, hashes"),
        ("Table1_space", "main Table 1: search spaces"),
        ("Table2_Q1_HV", "main Table 2: final hypervolume mean and 95% CI, verdicts"),
        ("Table3_speed", "main Table 3: median calls to 95% of the best hypervolume"),
        ("Table4_gap", "main Table 4: two ends of the pooled Pareto front per case"),
        ("Table5_designs", "main Table 5 and both front ends: designs, cell PCE, cost parts, CN/IN rate rows"),
        ("Fig2_HV_curves", "main Figure 2: mean and 95% band of hypervolume vs calls (and every seed)"),
        ("Fig3_front_points", "main Figure 3: Pareto front points per case (all feasible points are in Calls_all)"),
        ("Fig4_cost_parts", "main Figure 4: cost-part shares at the PCE-optimal design"),
        ("FigS1_bestPCE", "SI Figure S1: best module PCE so far vs calls (mean, band, every seed)"),
        ("FigS2_final_scores", "SI Figure S2 and Table S13: per-seed final scores and alternative verdicts"),
        ("FigS5_scans", "SI Figure S5 and Table S12: post-hoc gap scans"),
        ("TableS11_sensitivity", "SI Table S11: post-hoc gap under single changed assumptions"),
        ("FigS7_repeats", "SI Figure S7: repeated designs per run"),
        ("FigS8_P2_sensitivity", "SI Figure S8 and Table S8: earlier global sensitivity ranking"),
        ("Verification", "SI Table S9/S10: deterministic checks, re-execution and fresh re-simulation"),
        ("Physics_library", "every layer, interface and device parameter passed to SimSS, with provenance"),
        ("Cost_library", "every cost-library entry: value, unit, verbatim quotation, source, confidence"),
        ("Cost_sources", "registered source files with DOI, licence and SHA-256"),
        ("Cost_structure", "line inputs, operating-cost parts, materials bill decomposition (USD/m2)"),
        ("Assumptions", "every assumption, proxy and fixed setting of the study, with where it is used"),
        ("Numbers_main_text", "every number printed in the manuscript: token, printed value, raw value, source"),
    ]
    for l_ in lines:
        ws.append([l_])
    for n, d in sheets_doc:
        ws.append([f"  {n}", d])
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    # ------------------------------------------------------------------ Calls_all
    cols = ["arch", "scenario", "method", "seed", "call", "batch", "u (unit cube)", "x1_nm", "x2_nm", "x3_nm", "x4_nm",
            "stripe_mm", "layer names (x1..)", "design_hash", "sim_status", "reason", "sim_seconds", "cell_pce_pct", "voc_V",
            "ff", "jsc_mA_cm2", "gff", "p_tco", "module_pce_pct", "usd_per_W", "usd_per_W_high_prices", "materials_usd_m2",
            "cost_status", "feasible", "repeat_of_earlier_call_in_run"]
    rows = []
    for (a, s, m, k), r in sorted(runs.items()):
        names = [mm for _, mm in O.STACKS[a]]
        seen = set()
        for c in r["calls"]:
            th = c["x"][:-1] + [None] * (4 - len(c["x"][:-1]))
            rows.append([a, s, m, k, c["call"], c["batch"], c["u"], *th, c["x"][-1], "/".join(names), c["design_hash"],
                         c["sim_status"], c["reason"], c["seconds"], c["cell_pce_pct"], c["voc_V"], c["ff"], c["jsc_mA_cm2"],
                         c.get("gff"), c.get("p_tco"), c["module_pce_pct"], c["usd_per_W"], c.get("usd_per_W_high_prices"),
                         c.get("materials_usd_m2"), c.get("cost_status"), c["feasible"], c["design_hash"] in seen])
            seen.add(c["design_hash"])
    sheet(wb, "Calls_all", "All simulator calls (measured). One row per call; 90 runs x 240 calls.", cols, rows,
          notes=["Thickness columns follow the stack order in 'layer names'. p-i-n: NiOx/MAPbI3/C60/SnO2; n-i-p: C60/MAPbI3/PTAA "
                 "(x4 empty). Infeasible calls have no module PCE or cost. 'repeat' = same design hash asked earlier in the run "
                 "(cache hit, still one call)."],
          widths={"u (unit cube)": 40, "design_hash": 20, "reason": 30})

    # ------------------------------------------------------------------ Runs_summary
    rows = []
    for (a, s, m, k), r in sorted(runs.items()):
        mm = res["cases"][f"{a}/{s}"]["methods"][m]
        rows.append([a, s, m, k, r["hv_final"], r["n_infeasible"], mm["calls_to_95pct_best"][k],
                     max(c["module_pce_pct"] for c in r["calls"] if c["feasible"]),
                     min(c["usd_per_W"] for c in r["calls"] if c["feasible"]), r["seconds"], r["ref_point"][0], r["ref_point"][1],
                     r["sim_binary_sha256"], r["library_sha256"]])
    sheet(wb, "Runs_summary", "Run-level results (measured / computed).",
          ["arch", "scenario", "method", "seed", "hv_final", "n_infeasible", "calls_to_95pct_best (241 = not reached)",
           "best_module_pce_pct", "lowest_usd_per_W", "wall_seconds", "ref_pce_pct", "ref_usd_per_W", "sim_binary_sha256",
           "library_sha256"], rows)

    # ------------------------------------------------------------------ main tables
    rows = []
    for a in ARCHS:
        sp = O.space(a)
        for n, lo, hi in zip(sp.names, sp.lo, sp.hi):
            rows.append([a, n, lo, hi, "library thickness bound (phase0_bound)" if n != "stripe_mm" else "pre-registered"])
    sheet(wb, "Table1_space", "Main Table 1. Search spaces.", ["stack", "variable", "lower", "upper", "provenance"], rows)
    rows = []
    for k, c in res["cases"].items():
        for m in METHODS:
            mm = c["methods"][m]
            rows.append([k, m, mm["hv_mean"], mm["hv_ci95_half"], *mm["hv_final"],
                         ", ".join(kk for kk, b in c["q1"].items() if b)])
    sheet(wb, "Table2_Q1_HV", "Main Table 2. Final hypervolume after 240 calls (computed).",
          ["case", "method", "hv_mean", "hv_ci95_half (t, 4 dof)", "seed0", "seed1", "seed2", "seed3", "seed4", "verdicts (CI non-overlap)"], rows)
    rows = [[k, m, c["methods"][m]["calls_to_95pct_best_median"], *c["methods"][m]["calls_to_95pct_best"]]
            for k, c in res["cases"].items() for m in METHODS]
    sheet(wb, "Table3_speed", "Main Table 3. Calls to reach 95% of the best hypervolume of any run (241 = not reached).",
          ["case", "method", "median", "seed0", "seed1", "seed2", "seed3", "seed4"], rows)
    rows = []
    for k, c in res["cases"].items():
        q = c["q2"]
        rows.append([k, q["front_size"], q["n_feasible_pooled"], q["pce_optimal"]["module_pce_pct"], q["pce_optimal"]["usd_per_W"],
                     q["cost_optimal"]["module_pce_pct"], q["cost_optimal"]["usd_per_W"], q["gap_pce_points"], q["gap_cost_pct"], q["gap_small"]])
    sheet(wb, "Table4_gap", "Main Table 4. Two ends of the pooled Pareto front (computed).",
          ["case", "front_size", "n_feasible_pooled", "PCE-opt module_pce_pct", "PCE-opt usd_per_W", "cost-opt module_pce_pct",
           "cost-opt usd_per_W", "gap_pce_points", "gap_cost_pct", "gap_small (<1 pt and <5%)"], rows)
    rows = []
    for k, c in res["cases"].items():
        for end in ("pce_optimal", "cost_optimal"):
            d = c["q2"][end]
            rows.append([k, end, json.dumps(d["x"]), d["cell_pce_pct"], d["module_pce_pct"], d["usd_per_W"], d["materials_usd_m2"],
                         d["parts_usd_per_W"]["materials"], d["parts_usd_per_W"]["capital"], d["parts_usd_per_W"]["other_opex"],
                         d["usd_per_W_by_rate_region"]["MY"], d["usd_per_W_by_rate_region"]["CN"], d["usd_per_W_by_rate_region"]["IN"]])
    sheet(wb, "Table5_designs", "Main Table 5. Designs at both ends of each front, cost parts and capital-rate rows (computed).",
          ["case", "end", "design (nm, stripe mm)", "cell_pce_pct", "module_pce_pct", "usd_per_W", "materials_usd_m2",
           "part materials USD/W", "part capital USD/W", "part other opex USD/W", "USD/W MY rate", "USD/W CN rate", "USD/W IN rate"],
          rows, widths={"design (nm, stripe mm)": 60})

    # ------------------------------------------------------------------ figure data
    def ys(run):
        return [{"y": ([c["module_pce_pct"], c["usd_per_W"]] if c["feasible"] else None)} for c in run["calls"]]
    rows = []
    for a in ARCHS:
        for s in SCENS:
            ref = res["cases"][f"{a}/{s}"]["ref_point"]
            for m in METHODS:
                C = np.array([S.hv_curve(ys(runs[(a, s, m, k)]), ref) for k in range(5)])
                mu, sd = C.mean(0), C.std(0, ddof=1)
                for i in range(C.shape[1]):
                    rows.append([a, s, m, i + 1, mu[i], mu[i] - 2.776 * sd[i] / math.sqrt(5), mu[i] + 2.776 * sd[i] / math.sqrt(5), *C[:, i]])
    sheet(wb, "Fig2_HV_curves", "Main Figure 2. Hypervolume vs simulator calls (computed). Band = mean ± t(0.975,4) sd/sqrt(5).",
          ["arch", "scenario", "method", "call", "mean", "band_lo", "band_hi", "seed0", "seed1", "seed2", "seed3", "seed4"], rows)
    rows = []
    for a in ARCHS:
        for s in SCENS:
            cs = [c for (aa, ss, m, k), r in runs.items() if aa == a and ss == s for c in r["calls"] if c["feasible"]]
            front = S.pareto([(c["module_pce_pct"], c["usd_per_W"]) for c in cs])
            look = {(c["module_pce_pct"], c["usd_per_W"]): c for c in cs}
            for p, u in front:
                c = look[(p, u)]
                rows.append([a, s, p, u, c["materials_usd_m2"], json.dumps(c["x"]), c["cell_pce_pct"]])
    sheet(wb, "Fig3_front_points", "Main Figure 3. Pooled Pareto front points (stars). The scatter = all feasible rows of Calls_all.",
          ["arch", "scenario", "module_pce_pct", "usd_per_W", "materials_usd_m2", "design", "cell_pce_pct"], rows)
    rows = []
    for a in ARCHS:
        for s in SCENS:
            d = res["cases"][f"{a}/{s}"]["q2"]["pce_optimal"]
            t = d["usd_per_W"]
            rows.append([a, s, t, *(100 * d["parts_usd_per_W"][k] / t for k in ("materials", "capital", "other_opex"))])
    sheet(wb, "Fig4_cost_parts", "Main Figure 4. Cost-part shares at the PCE-optimal design (computed).",
          ["arch", "scenario", "total usd_per_W", "materials %", "capital %", "other opex %"], rows)
    rows = []
    for a in ARCHS:
        for s in SCENS:
            for m in METHODS:
                C = np.array([si["best_pce_curves"][f"{a}__{s}__{m}__s{k}"] for k in range(5)], float)
                mu, sd = C.mean(0), C.std(0, ddof=1)
                for i in range(C.shape[1]):
                    rows.append([a, s, m, i + 1, mu[i], mu[i] - 2.776 * sd[i] / math.sqrt(5), mu[i] + 2.776 * sd[i] / math.sqrt(5), *C[:, i]])
    sheet(wb, "FigS1_bestPCE", "SI Figure S1. Best module PCE found so far vs calls (computed).",
          ["arch", "scenario", "method", "call", "mean", "band_lo", "band_hi", "seed0", "seed1", "seed2", "seed3", "seed4"], rows)
    rows = []
    for k, c in si["q1_alternatives"].items():
        for m in METHODS:
            for metric, d in c["methods"][m].items():
                rows.append([k, m, metric, d["mean"], d["ci95_half"], *d["values"]])
        for metric, v in c["verdicts"].items():
            rows.append([k, "VERDICT", metric, ", ".join(v) or "none", None, None, None, None, None, None])
        rows.append([k, "REF", "tight ref point", json.dumps(c["ref_tight"]), None, None, None, None, None, None])
        rows.append([k, "REF", "nadir ref point", json.dumps(c["ref_nadir"]), None, None, None, None, None, None])
        rows.append([k, "RHO", "rho(tight HV, best PCE)", c["rho_hv_tight_vs_best_pce"], None, None, None, None, None, None])
    sheet(wb, "FigS2_final_scores", "SI Figure S2 / Table S13 (post-hoc). Per-seed final scores, alternative references and verdicts.",
          ["case", "method", "metric", "mean / value", "ci95_half", "seed0", "seed1", "seed2", "seed3", "seed4"], rows)
    rows = [[name, r["value"], r["arch"], r["scenario"], r["gap_pce_pt"], r["gap_cost_pct"], r["small"], r["front_size"],
             json.dumps(r["pce_opt_x"]), json.dumps(r["cost_opt_x"])] for name, rr in si["sensitivity"]["scans"].items() for r in rr]
    sheet(wb, "FigS5_scans", "SI Figure S5 / Table S12 (post-hoc re-scoring of stored designs).",
          ["scan", "value", "arch", "scenario", "gap_pce_pt", "gap_cost_pct", "small", "front_size", "PCE-opt design", "cost-opt design"], rows,
          notes=["absorber/all-layer price: multiplier on the per-nm layer price; throughput alpha: line cost per m2 x (1 + alpha "
                 "(t_abs/600 nm - 1)); R_sq: TCO sheet resistance (p_TCO scales linearly)."])
    rows = [[r["group"], r["change"], r["arch"], r["scenario"], r["front_size"], r["pce_opt_module_pct"], r["pce_opt_usd_W"],
             r["cost_opt_module_pct"], r["cost_opt_usd_W"], r["gap_pce_pt"], r["gap_cost_pct"], r["small"],
             json.dumps(r["x_pce_opt"]), json.dumps(r["x_cost_opt"]), r["n_designs"]] for r in si["sensitivity"]["rows"]]
    sheet(wb, "TableS11_sensitivity", "SI Table S11 (post-hoc). Gap after one changed assumption, all cases.",
          ["group", "change", "arch", "scenario", "front_size", "PCE-opt module %", "PCE-opt USD/W", "cost-opt module %",
           "cost-opt USD/W", "gap_pce_pt", "gap_cost_pct", "small", "PCE-opt design", "cost-opt design", "n designs"], rows)
    rows = []
    for (a, s, m, k), r in sorted(runs.items()):
        seen, n = set(), 0
        for c in r["calls"]:
            n += c["design_hash"] in seen
            seen.add(c["design_hash"])
        rows.append([a, s, m, k, n])
    sheet(wb, "FigS7_repeats", "SI Figure S7. Repeated designs per run (computed).", ["arch", "scenario", "method", "seed", "repeats"], rows)
    rows = [[k, i + 1, x["param"], x["index"], x["ci95_half"], x.get("S1")] for k in ("cell_pce_pct", "usd_per_W")
            for i, x in enumerate(p2["outputs"][k]["ranking"])]
    sheet(wb, "FigS8_P2_sensitivity", "SI Figure S8 / Table S8. Earlier global sensitivity (phase0/sensitivity/result.json).",
          ["output", "rank", "parameter", "Borgonovo delta", "ci95_half", "S1"], rows,
          notes=[f"{p2['n_unique_sims']} unique simulations; estimator: {p2.get('delta_estimator', '')}"[:250]])

    # ------------------------------------------------------------------ verification
    ver = si["verification"]
    rows = [["errors", json.dumps(ver["errors"]) or "{}"]] + [[f"worst {k}", v] for k, v in ver["worst"].items()]
    rows += [["shared_init_mismatch_groups", ver["shared_init_mismatch_groups"]], ["random_sobol_mismatch_runs", ver["random_sobol_mismatch_runs"]],
             ["binary_sha256", ", ".join(ver["binary_sha256"])], ["library_sha256", ", ".join(ver["library_sha256"])],
             ["n_calls", ver["n_calls"]], ["n_feasible", ver["n_feasible"]]]
    for p in sorted((A / "si").glob("rerun_*.json")):
        r = json.loads(p.read_text())
        rows.append([f"re-execution {r['run_id']}", f"identical designs {r['identical_x']}/{r['n_calls']}, identical objectives "
                     f"{r['identical_objectives']}/{r['n_calls']}, HV stored {r['hv_final_stored']} re-run {r['hv_final_rerun']}"])
    fp = A / "si" / "fresh_simss_check.json"
    if fp.exists():
        fr = json.loads(fp.read_text())
        rows.append(["fresh SimSS re-simulation (no cache)", f"n={fr['n']}, identical by field {json.dumps(fr['identical_by_field'])}"])
    sheet(wb, "Verification", "SI Note S8. Deterministic verification results (computed).", ["check", "result"], rows,
          widths={"check": 45, "result": 120})
    if fp.exists():
        fr = json.loads(fp.read_text())
        keys = list(fr["rows"][0].keys())
        sheet(wb, "Verification_fresh", "Fresh SimSS re-simulation of randomly chosen stored calls (no cache), stored vs new.",
              keys, [[r[k] for k in keys] for r in fr["rows"]])

    # ------------------------------------------------------------------ libraries
    rows = []
    for m, v in lib["materials"].items():
        for k, val in v["params"].items():
            rows.append(["layer", m, k, val, v["provenance"]["status"], v["provenance"].get("file"),
                         json.dumps(v.get("overridable", {}).get(k)) if k in v.get("overridable", {}) else None])
        rows.append(["layer", m, "thickness_nm (bounds)", f"{v['thickness_nm']['min']}-{v['thickness_nm']['max']}", v["thickness_nm"]["provenance"], None, None])
    for pair, v in lib["interfaces"].items():
        for k, val in v["params"].items():
            rows.append(["interface", pair, k, val, v["provenance"]["status"], v["provenance"].get("file") or v["provenance"].get("from"), None])
    for k, val in lib["device_setup"]["params"].items():
        rows.append(["device_setup", "all", k, val, lib["device_setup"]["provenance"]["status"], None, None])
    for a, ov in lib["architecture_setup"].items():
        for k, val in ov.items():
            rows.append(["architecture_setup", a, k, val, "override for this architecture", None, None])
    sheet(wb, "Physics_library", f"Physics library v{lib['library_version']} (SHA-256 {lib['_sha256']}). Units in key names (SI).",
          ["kind", "material / pair", "parameter", "value", "provenance status", "source file", "overridable range"], rows,
          notes=["simss_default = SIMsalabim v5.36 example values (not measured data); literature = read from a published table "
                 "(re-checked by tests); simss_default_transferred = ASSUMPTION; phase0_bound = design bound for convergence."])
    rows = []
    for e in clib["entries"]:
        rows.append([e["id"], e["value"], e.get("unit"), e.get("role"), e.get("item"), e["source"], e.get("quote_span"),
                     e.get("confidence"), e.get("cost_basis_year"), e.get("geography"), e.get("notes") or e.get("denominator_note")])
    sheet(wb, "Cost_library", f"Cost library v{clib['library_version']} entries (sourced; SHA-256 {clib['_sha256']}).",
          ["id", "value", "unit", "role", "item", "source", "verbatim quote", "confidence", "basis year", "geography", "notes"], rows,
          widths={"verbatim quote": 70, "notes": 50}, notes=["Each quote was found verbatim in the hashed source file by the provenance gate."])
    rows = [[k, s.get("citation"), s.get("doi"), s.get("licence"), s.get("geography"), s.get("file"), s.get("sha256"), s.get("notes")]
            for k, s in clib["sources"].items()]
    sheet(wb, "Cost_sources", "Registered cost sources.", ["key", "citation", "doi", "licence", "geography", "file", "sha256", "notes"], rows,
          widths={"citation": 60, "sha256": 66, "notes": 50})
    cs = si["cost_structure"]
    line = cs["line"]
    rows = [["line", k, line[k], line["tags"].get(k, "")] for k in ("capex_usd", "throughput_m2_per_yr", "depreciation_years", "rate", "yield_fraction")]
    rows += [["line", "crf", line["crf"], "computed: r(1+r)^n/((1+r)^n-1)"], ["line", "capital_usd_m2", line["capital_usd_m2"], "computed: capex*CRF/throughput"],
             ["line", "other_opex_usd_m2", line["other_opex_usd_m2"], line["tags"]["other_opex_usd_per_m2"]]]
    rows += [["other opex part", k, v, "computed (Note S4)"] for k, v in line["opex_parts_usd_m2"].items()]
    rows += [["other opex omitted", o, None, "not costed (lower bound)"] for o in line["opex_omitted"]]
    for a in ARCHS:
        rows += [[f"{a} fixed bill", k, v, "USD/m2, design-independent"] for k, v in cs[a]["fixed_items_usd_m2"].items()]
        rows += [[f"{a} layer slope", n, v["usd_m2_per_nm"], f"USD/m2 per nm (at reference {v['t_ref_nm']} nm: {v['usd_m2_at_ref']:.4f})"]
                 for n, v in cs[a]["layers"].items()]
    sheet(wb, "Cost_structure", "Cost structure used for every call (computed by the cost engine).", ["group", "item", "value", "tag / note"], rows,
          widths={"tag / note": 120})

    # ------------------------------------------------------------------ assumptions
    A_ = [
        ("module", "dead (interconnect) width", O.DEAD_MM, "mm", "assumed", "objective.DEAD_MM; same as earlier device comparison"),
        ("module", "front TCO sheet resistance", O.TCO_OHM_SQ, "ohm/sq", "assumed", "objective.TCO_OHM_SQ"),
        ("module", "edge-area fraction", O.EDGE, "fraction", "assumed", "objective.EDGE"),
        ("module", "max p_TCO accepted (first-order validity)", 0.15, "fraction", "model rule", "perodesign/module MAX_P_TCO"),
        ("device", "contacts ideal (S = infinite, semi-flat band work functions for p-i-n)", None, "", "assumed", "materials_v2.yaml device_setup"),
        ("device", "series resistance 0, shunt infinite", None, "", "assumed (SimSS defaults)", "device_setup"),
        ("device", "incident power", 1000, "W/m2", "standard (AM1.5G)", "perodesign/tea INCIDENT_W_M2"),
        ("device", "p-i-n interface parameters transferred from SimSS example interfaces", None, "", "assumed", "materials_v2.yaml interfaces"),
        ("device", "MAPbI3, C60, PTAA layer parameters = SimSS example values", None, "", "simss_default (not measured)", "materials_v2.yaml"),
        ("device", "thickness bounds", None, "nm", "phase0_bound (convergence)", "materials_v2.yaml thickness_nm"),
        ("scenario", "worse: N_t x sqrt(10), mu_n / sqrt(10)", O.SCENARIOS["worse"]["N_t_bulk_m3"], "m^-3", "convention", "objective.SCENARIOS"),
        ("scenario", "better: N_t / sqrt(10), mu_n x sqrt(10)", O.SCENARIOS["better"]["N_t_bulk_m3"], "m^-3", "convention", "objective.SCENARIOS"),
        ("cost", "PTAA density", O.ASM["density:PTAA"]["g_cm3"], "g/cm3", "assumed", "objective.ASM"),
        ("cost", "absorber priced as unspecified production perovskite (not MAPbI3)", None, "", "proxy", "cost_library layer_materials.MAPbI3"),
        ("cost", "front electrode priced as FTO glass (stack simulates ITO)", None, "", "proxy", "bom_per_m2 FTO_glass"),
        ("cost", "300 nm rear ITO costed but not simulated (p-i-n)", None, "", "simplification", "bom_by_architecture"),
        ("cost", "PTAA laboratory catalogue price; utilisation from slot-die perovskite", None, "", "proxy", "layer_materials.PTAA"),
        ("cost", "Au evaporation utilisation from Cu", None, "", "proxy", "back_electrode nk_Au.txt"),
        ("cost", "line capital: Chinese line, no regional adjustment", line["capex_usd"], "USD", "sourced, proxy for MY", "liu_capex_line"),
        ("cost", "capital rate = MY bank lending rate (lower bound)", line["rate"], "fraction", "sourced, proxy", "wdi_lend_my_2025"),
        ("cost", "throughput independent of thickness", line["throughput_m2_per_yr"], "m2/yr", "assumed (tested post-hoc, Note S9)", "liu2025_line"),
        ("cost", "other opex excludes land/rent, HVAC base load, employer contributions (service tax not applicable: SI Note S17)", line["other_opex_usd_m2"], "USD/m2", "lower bound", "region_line_MY_bottom_up"),
        ("cost", "staffing from Cai 2017 printable line; kWh/m2 from Tian 2020 tandem inventory; maintenance 1%/yr (Kajal)", None, "", "proxies", "tea module"),
        ("optimiser", "budget per run", S.BUDGET, "calls", "pre-registered", "search.BUDGET"),
        ("optimiser", "shared Sobol initial design", S.N_INIT, "points", "pre-registered", "search.N_INIT"),
        ("optimiser", "qNEHVI batch size / restarts / raw samples / MC samples", "12 / 4 / 256 / 64", "", "pre-registered", "search.run_qnehvi"),
        ("optimiser", "NSGA-II population", S.POP, "", "pre-registered", "search.POP"),
        ("metric", "HV reference point module PCE", O.PCE_FLOOR, "%", "pre-registered", "objective.PCE_FLOOR"),
        ("metric", "HV reference cost p-i-n / n-i-p", f"{O.reference_point('p-i-n')[1]:.4f} / {O.reference_point('n-i-p')[1]:.4f}", "USD/W",
         "pre-registered rule (max-thickness design at PCE floor)", "objective.reference_point"),
        ("rule", "small gap", "dPCE < 1 point and dCost < 5%", "", "pre-registered", "docs/PaperA_preregistration.md s7"),
        ("rule", "design robustness", "absorber moves < 20% of range", "", "pre-registered", "s7"),
        ("rule", "unreliable case", "> 12 infeasible calls in one run", "", "pre-registered", "s7"),
        ("rounding", "thickness / stripe rounding", "0.1 nm / 0.01 mm", "", "pre-registered", "objective.design"),
    ]
    sheet(wb, "Assumptions", "Every assumption, proxy and fixed setting of the study.",
          ["area", "assumption / setting", "value", "unit", "status", "where in code / library"], [list(a) for a in A_],
          widths={"assumption / setting": 80, "where in code / library": 45})

    rows = [[k, v["value"], v["raw"] if not isinstance(v["raw"], (list, dict)) else json.dumps(v["raw"]), v["source"], "numbers.json"] for k, v in N.items()]
    rows += [[k, v["value"], v["raw"], v["source"], "numbers_si.json (post-hoc)"] for k, v in NS.items()]
    sheet(wb, "Numbers_main_text", "Every token the manuscript can print (value as printed, raw value, source).",
          ["token", "printed", "raw", "source", "file"], rows, widths={"source": 100})

    if MERGED:
        merged_numbering(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    # ------------------------------------------------------------------ read back and check
    wb2 = load_workbook(OUT, read_only=True)
    ws = wb2["Calls_all"]
    data = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(data) if r and r[0] == "arch")
    body = data[hdr + 1:]
    assert len(body) == 21600, len(body)
    idx = {c: i for i, c in enumerate(data[hdr])}
    import random
    for row in random.Random(1).sample(body, 500):
        c = runs[(row[idx["arch"]], row[idx["scenario"]], row[idx["method"]], row[idx["seed"]])]["calls"][row[idx["call"]]]
        for key in ("usd_per_W", "module_pce_pct", "cell_pce_pct", "materials_usd_m2"):
            a_, b_ = row[idx[key]], c.get(key)
            assert (a_ is None and b_ is None) or math.isclose(a_, b_, rel_tol=1e-12), (key, a_, b_)
    feas = sum(1 for r in body if r[idx["feasible"]])
    print(json.dumps({"xlsx": str(OUT), "bytes": OUT.stat().st_size, "sheets": wb2.sheetnames, "calls_rows": len(body),
                      "feasible_rows": feas, "sample_check": "500/500 rows match run files"}))


if __name__ == "__main__":
    main()
